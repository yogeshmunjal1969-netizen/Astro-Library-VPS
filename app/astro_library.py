#!/usr/bin/env python3
"""
astro_library.py - Catalogue, search and organise the downloaded astrology library.

Works on the folder created by astro_downloader.py (default F:\\AstroBooksCollection).
All data lives in  <library>\\_library\\library.db  (SQLite + FTS5 full-text index).

Metadata is derived from
  * the collector's file-naming convention, e.g.
        Book_1991_Jeanne Avery_The Rising Sign_Your Astrological Mask.pdf
        Jyotish_2021_B Hariharan_Matching of Horoscopes for Marriages_must read.pdf
        Russian_<title>_<author>.pdf
  * the folder hierarchy  (topic / sub-topic / magazine series / language collection)
  * embedded PDF metadata (title, author, publisher, page count) where available
  * a dictionary of well-known astrology publishers

Commands
  build                      (re)build the catalogue  [--no-pdf-meta to skip reading PDFs]
  search [text] [filters]    search; filters: --author --title --topic --publisher --year
                             --from --to --lang --type --tradition --tag --ext  (--open N opens hit N)
  authors|topics|publishers|years|languages|types [--like X]   list values with counts
  stats                      collection overview
  open <id>                  open a book by its catalogue id
  export --csv FILE | --xlsx FILE   export the catalogue
  html                       regenerate the offline catalogue  "<library>\\Astro Library Catalogue.html"
  views [--by author,year,topic,publisher,language,type]
                             create browse folders under <library>\\_Browse using NTFS hard links
                             (no extra disk space is used)

Examples
  python astro_library.py build
  python astro_library.py search navamsa
  python astro_library.py search --author "Alan Leo" --from 1890 --to 1920
  python astro_library.py search "medical astrology" --type Book --tag "must read"
  python astro_library.py search --publisher Sagar --open 1
  python astro_library.py authors --like raman
"""

import argparse
import csv
import html
import json
import os
import re
import sqlite3
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime

DEFAULT_LIB = os.environ.get("ASTRO_LIB") or r"F:\AstroBooksCollection"
DOC_EXTS = {"pdf", "doc", "docx", "rtf", "txt", "epub", "mobi", "djvu", "htm", "html", "webarchive", "odt", "pps",
            "ppt", "pptx", "xls", "xlsx", "mp4", "mp3", "chm", "azw3"}

# ----------------------------------------------------------------------------- vocabularies
PREFIX_TYPES = {
    "book": ("Book", None, None), "books": ("Book", None, None),
    "jyotish": ("Book", "Vedic / Jyotish", None),
    "article": ("Article", None, None), "articles": ("Article", None, None),
    "paper": ("Article", None, None), "thesis": ("Thesis", None, None),
    "presentation": ("Presentation", None, None), "magazine": ("Magazine", None, None),
    "journal": ("Magazine", None, None), "course": ("Course / Lessons", None, None),
    "lesson": ("Course / Lessons", None, None), "lessons": ("Course / Lessons", None, None),
    "notes": ("Notes", None, None), "ebook": ("Book", None, None),
    "russian": (None, None, "Russian"), "french book": ("Book", None, "French"), "french": (None, None, "French"),
    "german": (None, None, "German"), "spanish": (None, None, "Spanish"), "italian": (None, None, "Italian"),
    "hindi": (None, None, "Hindi"), "tamil": (None, None, "Tamil"), "kp": (None, "KP (Krishnamurti Paddhati)", None),
    "kp ezine": ("Magazine", "KP (Krishnamurti Paddhati)", None), "nadi": (None, "Vedic / Jyotish", None),
}
TAG_WORDS = {
    "kindle": "kindle", "must read": "must read", "useless": "useless", "retyped": "retyped",
    "scanned": "scanned", "ocr": "ocr", "rare": "rare", "translated": "translated", "draft": "draft",
    "incomplete": "incomplete", "summary": "summary", "notes": "notes", "excerpt": "excerpt",
    "good": "good", "very good": "very good", "excellent": "excellent", "best": "best", "old": "old",
    "dual language": "dual language", "bilingual": "dual language", "ai translation": "ai translation",
    "free": "free", "sample": "sample", "preview": "preview", "abridged": "abridged", "reprint": "reprint",
}
LANGS = ["English", "Hindi", "Sanskrit", "Tamil", "Telugu", "Kannada", "Malayalam", "Marathi", "Bengali",
         "Gujarati", "Oriya", "Odia", "Punjabi", "Urdu", "Nepali", "Russian", "French", "German", "Spanish",
         "Italian", "Portuguese", "Dutch", "Polish", "Czech", "Persian", "Arabic", "Latin", "Greek", "Chinese",
         "Japanese", "Turkish", "Hungarian", "Romanian", "Bulgarian", "Serbian", "Croatian", "Ukrainian"]
PUBLISHERS = [
    "Sagar Publications", "Motilal Banarsidass", "Ranjan Publications", "Sterling Publishers", "Sterling",
    "Llewellyn", "Penguin", "Hay House", "American Federation of Astrologers", "AFA", "Weiser", "Samuel Weiser",
    "Wessex Astrologer", "Astrology Center of America", "Lotus Press", "Chowkhamba", "Chaukhamba", "Chaukhambha",
    "Krishnamurti Publications", "Raman Publications", "IBH Prakashana", "Pustak Mahal", "Manoj Publications",
    "Alpha Publications", "Upkar Prakashan", "Diamond Pocket Books", "Diamond Books", "Rupa", "Bharatiya Vidya Bhavan",
    "Gita Press", "Khemraj", "Venkateshwar Press", "Nag Publishers", "Parimal Publications", "Kalpataru",
    "Bantam", "Ballantine", "Arkana", "Routledge", "Springer", "Cambridge University Press", "Oxford University Press",
    "HarperCollins", "Harper & Row", "Harper", "Simon & Schuster", "Red Wheel", "Inner Traditions", "Destiny Books",
    "CRCS", "Aquarian Press", "Thorsons", "Element Books", "Flare Publications", "Astrolabe", "Project Hindsight",
    "Arhat", "Sri Satguru Publications", "Indological Book House", "Munshiram Manoharlal", "Bhartiya Vidya Prakashan",
    "Saptarishis Astrology", "Saptarishis", "Vedic Astrology Research Portal", "Astrological Magazine",
    "Bangalore Raman", "Brihaspati Publications", "Shri Ram Publications", "Kala Prakashan", "Jaico",
    "Theosophical Publishing House", "Theosophical", "Rider", "Dover", "Wiley", "Fowler", "L.N. Fowler",
    "Regulus Publishing", "Ascella", "Astro Computing Services", "ACS Publications", "Starcrafts",
    "Kosmos", "Ebertin Verlag", "Penguin Arkana", "Lulu", "CreateSpace", "Notion Press", "Partridge",
    "Blue Rose", "BookLeaf", "Pothi", "Leadstart", "Prabhat Prakashan", "Rajpal", "Vani Prakashan",
    "Shubhi Publications", "Sanskriti", "Hind Pocket Books", "Orient Paperbacks", "Full Circle", "New Age Books",
    "Pilgrims Publishing", "Bharatiya Kala Prakashan", "Eastern Book Linkers", "Nirnaya Sagar",
]
TITLE_WORDS = {
    "the", "a", "an", "of", "to", "in", "on", "for", "with", "your", "you", "my", "how", "what", "why", "who",
    "is", "are", "at", "by", "from", "into", "about", "through", "astrology", "astrological", "astrologer",
    "horoscope", "horoscopes", "guide", "book", "books", "secrets", "secret", "introduction", "handbook",
    "complete", "art", "science", "vol", "volume", "part", "magazine", "journal", "vedic", "jyotish", "planet",
    "planets", "signs", "sign", "houses", "house", "love", "life", "moon", "sun", "mars", "venus", "mercury",
    "jupiter", "saturn", "rahu", "ketu", "chart", "charts", "zodiac", "stars", "star", "dasa", "dasha", "yoga",
    "yogas", "medical", "health", "marriage", "career", "predictive", "prediction", "predictions", "astronomy",
    "lessons", "lesson", "course", "study", "studies", "principles", "practical", "ancient", "modern", "new",
    "essentials", "beginners", "advanced", "theory", "system", "systems", "method", "methods", "nakshatra",
    "nakshatras", "transits", "transit", "aspects", "ascendant", "remedies", "tarot", "numerology", "palmistry",
    "kp", "nadi", "jaimini", "horary", "prasna", "muhurta", "varga", "vargas", "divisional", "navamsa", "and",
    "edition", "ed", "issue", "series", "notes", "tables", "ephemeris", "manual", "encyclopedia", "dictionary",
    "astro", "cosmic", "karma", "soul", "spiritual", "western", "indian", "hindu", "sidereal", "tropical",
    "kindle", "useless", "retyped", "article", "articles", "research", "report", "seminar", "conference",
    "unknown", "anonymous", "various", "misc", "collection", "compilation", "copy", "untitled", "english",
    "hindi", "tamil", "sanskrit", "russian", "french", "translation", "translated", "dual", "language",
}
NAME_CONNECTORS = {"and", "&", "de", "van", "von", "du", "da", "di", "le", "la", "der", "den", "bin", "ben", "y", "del"}
YEAR_RE = re.compile(r"^\(?(-?\d{3,4})(?:\s*(?:-|to)\s*-?\d{1,4})?\s*(BCE|BC|AD|CE)?\)?\.?$", re.I)


_PUB_LONG = re.compile("(?<![A-Za-z])(" + "|".join(re.escape(p) for p in sorted(PUBLISHERS, key=len, reverse=True)
                                                   if len(p) > 4) + ")(?![A-Za-z])", re.I)
_PUB_SHORT = re.compile("(?<![A-Za-z])(" + "|".join(re.escape(p) for p in PUBLISHERS if len(p) <= 4) + ")(?![A-Za-z])")
_PUB_CANON = {p.lower(): p for p in PUBLISHERS}


def find_publisher(text):
    if not text:
        return None
    m = _PUB_LONG.search(text) or _PUB_SHORT.search(text)
    return _PUB_CANON.get(m.group(1).lower(), m.group(1)) if m else None


def clean(s):
    return re.sub(r"\s+", " ", s.replace("_", " ")).strip(" -.,;")


def pretty_folder(name):
    """'1_PLANETS_Navagrahas_MAIN ONES' -> 'PLANETS, Navagrahas, MAIN ONES'"""
    name = re.sub(r"^\d{1,2}_", "", name)                 # only list-ordering prefixes like "1_" or "6_"
    parts = [p.strip() for p in name.split("_") if p.strip()]
    return ", ".join(parts) if parts else name


def norm_key(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode() or s
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def looks_like_name(seg, known=None):
    seg = seg.strip()
    if not seg or len(seg) > 60 or re.search(r"\d", seg):
        return False
    if known and norm_key(seg) in known:
        return True
    if "," in seg:                                         # "Merrick, Richard"
        a, _, b = seg.partition(",")
        return looks_like_name(a) and looks_like_name(b) if b.strip() and " " not in a.strip() else False
    words = re.split(r"\s+", seg.replace(".", ". ").strip())
    words = [w for w in words if w]
    real = [w for w in words if w.lower() not in NAME_CONNECTORS]
    if not 2 <= len(real) <= 6:
        return False
    for w in real:
        core = w.strip(".()'’-")
        if not core:
            continue
        if core.lower() in TITLE_WORDS:
            return False
        if not (core[0].isupper() or (len(core) <= 2)):
            return False
    return True


def fix_name(seg):
    seg = clean(seg)
    if "," in seg:
        last, _, first = seg.partition(",")
        if first.strip() and " " not in last.strip():
            seg = f"{first.strip()} {last.strip()}"
    return re.sub(r"\s+", " ", seg)


def split_authors(s):
    parts = re.split(r"\s+(?:and|&)\s+|;\s*", s)
    return [p.strip() for p in parts if p.strip()]


# ----------------------------------------------------------------------------- parsing
def parse_record(row, known_authors=None, folder_authors=None):
    """row: dict with remote_path, name, size ... -> metadata dict"""
    rp = row["remote_path"]
    parts = rp.split("/")
    folders, fname = parts[:-1], parts[-1]
    stem, ext = os.path.splitext(fname)
    ext = ext.lower().lstrip(".")
    rec = dict(title=None, authors=None, year=None, year_text=None, publisher=None, doc_type=None,
               tradition=None, language=None, tags=[], ext=ext, collection=folders[0] if folders else "(root)",
               topic=None, subtopic=None, confident=False)

    # ---------- topic from folders
    col = rec["collection"]
    if col.lower().startswith("classified by topics") and len(folders) > 1:
        rec["topic"] = pretty_folder(folders[1])
        rec["subtopic"] = " › ".join(pretty_folder(f) for f in folders[2:]) or None
    else:
        rec["topic"] = pretty_folder(col)
        rec["subtopic"] = " › ".join(pretty_folder(f) for f in folders[1:]) or None
    path_low = rp.lower()

    # ---------- collection-level defaults
    if col.lower().startswith("magazine"):
        rec["doc_type"] = "Magazine"
        if len(folders) > 1:
            rec["publisher"] = pretty_folder(folders[1])            # magazine series acts as publisher
    if "russian" in col.lower():
        rec["language"] = "Russian"
    if "indian languages" in col.lower():
        rec["language"] = "Indian language"
    if col.lower().startswith("kp "):
        rec["tradition"] = "KP (Krishnamurti Paddhati)"
    if col.lower().startswith("my compilation"):
        rec["doc_type"] = "Notes / Compilation"
    if "youtube" in col.lower():
        rec["doc_type"] = "Video / Transcript"

    # ---------- split file name
    segs = [s.strip() for s in stem.split("_")]
    segs = [s for s in segs if s]

    # trailing tags
    while segs:
        last = segs[-1].strip().lower().strip("()")
        if last in TAG_WORDS:
            rec["tags"].append(TAG_WORDS[last]); segs.pop()
        else:
            break
    # tags embedded at end of title e.g. "... 2026_kindle" handled above; "(dual Hindi + English)" -> language
    m = re.search(r"\((?:dual\s+)?([A-Za-z]+)\s*\+\s*([A-Za-z]+)\)|\(dual language\)", stem, re.I)
    if m:
        rec["tags"].append("dual language")
        if m.group(1):
            rec["language"] = f"{m.group(1).title()} + {m.group(2).title()}"

    # "2013-Maria Cristina Casabur" -> "2013", "Maria Cristina Casabur"
    for i in range(min(3, len(segs))):
        ym = re.match(r"^(\d{4})\s*-\s*([A-Za-z].*)$", segs[i])
        if ym:
            segs[i:i + 1] = [ym.group(1), ym.group(2)]
            break

    # leading prefixes (Book_, Jyotish_, Russian_, KP_ ...)
    for _ in range(3):
        if not segs:
            break
        k = segs[0].strip().lower()
        if k in PREFIX_TYPES:
            dt, trad, lang = PREFIX_TYPES[k]
            rec["doc_type"] = dt or rec["doc_type"]
            rec["tradition"] = trad or rec["tradition"]
            rec["language"] = lang or rec["language"]
            rec["confident"] = rec["confident"] or k in ("book", "jyotish", "french book", "article")
            segs.pop(0)
        elif re.match(r"^article\s*\d{4}", k) or k.startswith("article"):
            rec["doc_type"] = "Article"; break
        else:
            break

    # year (first or second segment, or "-1400-1200 BCE_article_...")
    for i in range(min(2, len(segs))):
        ym = YEAR_RE.match(segs[i])
        if ym:
            y = int(ym.group(1))
            if ym.group(2) and ym.group(2).upper() in ("BC", "BCE"):
                y = -abs(y)
            if -3500 <= y <= datetime.now().year + 1:
                rec["year"], rec["year_text"] = y, segs[i].strip("(). ")
                segs.pop(i)
                if i == 0:
                    rec["confident"] = True
                # "1933_article_..." -> another type word may follow the year
                if segs and segs[0].lower() in PREFIX_TYPES:
                    dt, trad, lang = PREFIX_TYPES[segs[0].lower()]
                    rec["doc_type"] = dt or rec["doc_type"]; rec["tradition"] = trad or rec["tradition"]
                    rec["language"] = lang or rec["language"]; segs.pop(0)
                break
    if rec["year"] is None:
        for s in reversed(segs):                            # "... Guide 2026" or "(1998)"
            ym = re.search(r"(?<![\d.])(1[5-9]\d\d|20[0-3]\d)(?![\d])", s)
            if ym and (s.strip().endswith(ym.group(1)) or s.strip().startswith(ym.group(1)) or "(" + ym.group(1) in s):
                rec["year"], rec["year_text"] = int(ym.group(1)), ym.group(1)
                break

    # author
    author = None
    # files without a Book_/Jyotish_/year prefix only accept authors already seen elsewhere
    strict = None if (rec["confident"] or rec["year"] is not None) else (known_authors or set())

    def name_ok(seg):
        return norm_key(seg) in strict if strict is not None else looks_like_name(seg, known_authors)

    def name_of(seg):                                    # "Abha Sharma (Guide K N Rao)" -> "Abha Sharma"
        return re.sub(r"\s*\([^)]*\)\s*", " ", seg).strip()

    if segs and len(segs) >= 2 and name_ok(name_of(segs[0])):
        author = name_of(segs.pop(0))
    elif segs and len(segs) >= 2 and name_ok(name_of(segs[-1])):
        author = name_of(segs.pop())
    elif known_authors and segs:
        for i, s in enumerate(segs):
            if norm_key(s) in known_authors and len(segs) > 1:
                author = segs.pop(i); break
    if not author and folder_authors:
        for f in reversed(folders):
            if f in folder_authors:
                author = folder_authors[f]; break
    if not author:
        for f in reversed(folders):                         # folder "by K.S. Krishnamurti (K Subramaniam)"
            fm = re.match(r"^by\s+(.+?)(?:\s*\(.*\))?$", f.strip(), re.I)
            if fm:
                author = fm.group(1); break
    if author:
        rec["authors"] = "; ".join(fix_name(a) for a in split_authors(author))

    # title
    if segs:
        title = clean(segs[0])
        if len(segs) > 1:
            title += ": " + " - ".join(clean(s) for s in segs[1:])
        rec["title"] = title or clean(stem) or fname
    else:
        rec["title"] = clean(stem) or fname

    # ---------- language / tradition / type inference
    if not rec["language"] or rec["language"] == "Indian language":
        found = None
        hits = [lang for lang in LANGS if re.search(rf"(?<![A-Za-z]){lang}(?![A-Za-z])", stem, re.I)]
        if len(hits) > 1 and "English" in hits:
            found = " + ".join([h for h in hits if h != "English"][:1] + ["English"])
        elif hits:
            found = hits[0]
        if re.search(r"[\u0400-\u04FF]", stem):
            found = found or "Russian"
        elif re.search(r"[\u0900-\u097F]", stem):
            found = found or "Hindi / Sanskrit"
        elif re.search(r"[\u0B80-\u0BFF]", stem):
            found = found or "Tamil"
        rec["language"] = found or rec["language"] or "English"
    if not rec["tradition"]:
        if re.search(r"\bkp\b|krishnamurti", path_low):
            rec["tradition"] = "KP (Krishnamurti Paddhati)"
        elif re.search(r"vedic|jyotish|jaimini|nadi|parasara|parashara|nakshatra|dasa|dasha|varga|navamsa|panchang|muhurt|prasna|tajik|lal kitab|bhrigu", path_low):
            rec["tradition"] = "Vedic / Jyotish"
        elif re.search(r"western|tropical|medieval|hellenistic|uranian|horary|evangeline|alan leo|rudhyar", path_low):
            rec["tradition"] = "Western"
        elif rec["doc_type"] == "Book" and (segs or True) and stem.lower().startswith("book_"):
            rec["tradition"] = "Western"
        elif "chinese" in path_low:
            rec["tradition"] = "Chinese"
        else:
            rec["tradition"] = "General / Mixed"
    if not rec["doc_type"]:
        if ext in ("htm", "html", "webarchive"):
            rec["doc_type"] = "Web page"
        elif ext in ("jpg", "jpeg", "png", "gif", "bmp", "tif", "tiff"):
            rec["doc_type"] = "Image"
        elif ext in ("mp4", "mp3", "mov", "avi", "m4a"):
            rec["doc_type"] = "Video / Audio"
        elif "article" in path_low or "journal" in path_low or "academia" in path_low:
            rec["doc_type"] = "Article"
        elif ext in ("doc", "docx", "rtf", "txt"):
            rec["doc_type"] = "Document"
        else:
            rec["doc_type"] = "Book / Document"
    if "kindle" in rec["tags"] and not rec["publisher"]:
        rec["publisher"] = "Kindle edition"

    # publisher from file name / folders
    p = find_publisher(rp)
    if p and rec["publisher"] in (None, "Kindle edition"):
        rec["publisher"] = p
    rec["tags"] = sorted(set(rec["tags"]))
    return rec


# ----------------------------------------------------------------------------- PDF metadata
def read_pdf_meta(path):
    try:
        import pypdf
    except ImportError:
        return None
    import logging
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    try:
        r = pypdf.PdfReader(path, strict=False)
        info = r.metadata or {}
        meta = {
            "pdf_title": str(info.get("/Title") or "").strip() or None,
            "pdf_author": str(info.get("/Author") or "").strip() or None,
            "pdf_subject": str(info.get("/Subject") or "").strip() or None,
            "pdf_keywords": str(info.get("/Keywords") or "").strip() or None,
            "pdf_publisher": str(info.get("/Publisher") or info.get("/Company") or "").strip() or None,
            "pages": None,
        }
        try:
            meta["pages"] = len(r.pages)
        except Exception:
            pass
        return meta
    except Exception:
        return {"pages": None}


JUNK_META = re.compile(r"^(untitled|microsoft word|document\d*|unknown|admin|user|owner|pdf|scan|\W*|.*\.(docx?|pdf|indd|qxd))$", re.I)


def meta_ok(v):
    return bool(v) and len(v) < 200 and not JUNK_META.match(v.strip())


# ----------------------------------------------------------------------------- database
SCHEMA = """
CREATE TABLE IF NOT EXISTS pdfmeta (
    fileid INTEGER PRIMARY KEY, hash TEXT, pages INTEGER, pdf_title TEXT, pdf_author TEXT,
    pdf_subject TEXT, pdf_keywords TEXT, pdf_publisher TEXT);
DROP TABLE IF EXISTS books;
CREATE TABLE books (
    id INTEGER PRIMARY KEY, fileid INTEGER UNIQUE, title TEXT, authors TEXT, year INTEGER, year_text TEXT,
    publisher TEXT, doc_type TEXT, tradition TEXT, language TEXT, tags TEXT, ext TEXT,
    collection TEXT, topic TEXT, subtopic TEXT, pages INTEGER, size INTEGER,
    local_path TEXT, remote_path TEXT, downloaded INTEGER, added TEXT);
CREATE INDEX ix_books_year ON books(year);
CREATE INDEX ix_books_topic ON books(topic);
CREATE INDEX ix_books_authors ON books(authors);
DROP TABLE IF EXISTS books_fts;
CREATE VIRTUAL TABLE books_fts USING fts5(
    title, authors, publisher, topic, subtopic, tags, language, tradition, doc_type, filename,
    content='', tokenize='unicode61 remove_diacritics 2');
"""


def open_db(lib):
    path = os.path.join(lib, "_library", "library.db")
    if not os.path.exists(path):
        sys.exit(f"No manifest at {path}. Run astro_downloader.py first.")
    db = sqlite3.connect(path)
    db.row_factory = sqlite3.Row
    return db


def long_path(p):
    if os.name == "nt":
        p = os.path.abspath(p)
        if p.startswith("\\\\?\\"):
            pass
        elif p.startswith("\\\\"):                    # network share: \\server\share\...
            p = "\\\\?\\UNC\\" + p[2:]
        else:
            p = "\\\\?\\" + p
    return p


def cmd_build(db, lib, args):
    db.executescript(SCHEMA)
    rows = [dict(r) for r in db.execute(
        "SELECT * FROM files WHERE status NOT IN ('skipped','removed') ORDER BY remote_path")]
    rows = [r for r in rows if os.path.splitext(r["name"])[1].lower().lstrip(".") in DOC_EXTS]
    print(f"Cataloguing {len(rows)} documents ...")

    # pass 1: confident parses -> dictionary of known authors
    first = [parse_record(r) for r in rows]
    counts = Counter()
    display = defaultdict(Counter)
    for rec in first:
        if rec["authors"] and rec["confident"]:
            for a in rec["authors"].split("; "):
                counts[norm_key(a)] += 1
                display[norm_key(a)][a] += 1
    known = {k for k, n in counts.items() if n >= 2 and len(k) > 3}
    # folders that are clearly named after an author
    folder_authors = {}
    for r in rows:
        for f in r["remote_path"].split("/")[:-1]:
            if norm_key(f) in known:
                folder_authors[f] = f
    print(f"  {len(known)} authors learnt from file names")

    # PDF metadata (incremental: only files not yet read or changed)
    if not args.no_pdf_meta:
        have = {r[0]: r[1] for r in db.execute("SELECT fileid, hash FROM pdfmeta")}
        todo = [r for r in rows if r["status"] == "done" and r["name"].lower().endswith(".pdf")
                and have.get(r["fileid"]) != r["hash"] and (r["size"] or 0) < 400 * 1024 * 1024]
        if todo:
            print(f"  reading embedded metadata from {len(todo)} PDFs (first run is slow; later runs are incremental)")
        for i, r in enumerate(todo, 1):
            m = read_pdf_meta(long_path(os.path.join(lib, r["local_path"])))
            if m is not None:
                db.execute("INSERT OR REPLACE INTO pdfmeta VALUES (?,?,?,?,?,?,?,?)",
                           (r["fileid"], r["hash"], m.get("pages"), m.get("pdf_title"), m.get("pdf_author"),
                            m.get("pdf_subject"), m.get("pdf_keywords"), m.get("pdf_publisher")))
            if i % 250 == 0:
                db.commit()
                print(f"    {i}/{len(todo)}", flush=True)
        db.commit()
    pdfmeta = {r["fileid"]: dict(r) for r in db.execute("SELECT * FROM pdfmeta")}

    # pass 2: final records
    canon = {k: display[k].most_common(1)[0][0] for k in display}
    n = 0
    for r in rows:
        rec = parse_record(r, known, folder_authors)
        pm = pdfmeta.get(r["fileid"], {})
        if not rec["authors"] and meta_ok(pm.get("pdf_author")) and looks_like_name(pm["pdf_author"].split(";")[0]):
            rec["authors"] = pm["pdf_author"]
        if pm.get("pdf_publisher") and meta_ok(pm["pdf_publisher"]) and rec["publisher"] in (None, "Kindle edition"):
            rec["publisher"] = pm["pdf_publisher"]
        if not rec["publisher"]:
            rec["publisher"] = find_publisher(" ".join(filter(None, [pm.get("pdf_subject"), pm.get("pdf_keywords"),
                                                                     pm.get("pdf_title")])))
        if rec["authors"]:
            rec["authors"] = "; ".join(canon.get(norm_key(a), a) for a in rec["authors"].split("; "))
        tags = ", ".join(rec["tags"])
        cur = db.execute(
            "INSERT INTO books(fileid,title,authors,year,year_text,publisher,doc_type,tradition,language,tags,ext,"
            "collection,topic,subtopic,pages,size,local_path,remote_path,downloaded,added) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (r["fileid"], rec["title"], rec["authors"], rec["year"], rec["year_text"], rec["publisher"],
             rec["doc_type"], rec["tradition"], rec["language"], tags, rec["ext"], rec["collection"],
             rec["topic"], rec["subtopic"], pm.get("pages"), r["size"], r["local_path"], r["remote_path"],
             1 if r["status"] == "done" else 0, r["created"]))
        db.execute("INSERT INTO books_fts(rowid,title,authors,publisher,topic,subtopic,tags,language,tradition,"
                   "doc_type,filename) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                   (cur.lastrowid, rec["title"], rec["authors"], rec["publisher"], rec["topic"], rec["subtopic"],
                    tags, rec["language"], rec["tradition"], rec["doc_type"], r["remote_path"].replace("_", " ")))
        n += 1
    db.commit()
    print(f"Catalogue built: {n} entries.")
    cmd_stats(db, lib, args)
    cmd_html(db, lib, args)


# ----------------------------------------------------------------------------- queries
def fts_query(text):
    terms = re.findall(r'"[^"]+"|\S+', text)
    out = []
    for t in terms:
        if t.startswith('"'):
            out.append(t)
        else:
            t = re.sub(r'[^\w\u0080-\uFFFF-]', "", t)
            if t:
                out.append(f'"{t}"*')
    return " AND ".join(out)


def search(db, a):
    where, params = [], []
    if a.text:
        q = fts_query(" ".join(a.text))
        if q:
            where.append("b.id IN (SELECT rowid FROM books_fts WHERE books_fts MATCH ?)")
            params.append(q)
    for col, val in (("authors", a.author), ("title", a.title), ("publisher", a.publisher),
                     ("language", a.lang), ("doc_type", a.type), ("tradition", a.tradition), ("tags", a.tag)):
        if val:
            where.append(f"b.{col} LIKE ?"); params.append(f"%{val}%")
    if a.topic:
        where.append("(b.topic LIKE ? OR b.subtopic LIKE ?)"); params += [f"%{a.topic}%"] * 2
    if a.ext:
        where.append("b.ext = ?"); params.append(a.ext.lower().lstrip("."))
    if a.year is not None:
        where.append("b.year = ?"); params.append(a.year)
    if a.year_from is not None:
        where.append("b.year >= ?"); params.append(a.year_from)
    if a.year_to is not None:
        where.append("b.year <= ?"); params.append(a.year_to)
    if a.downloaded:
        where.append("b.downloaded = 1")
    order = {"year": "b.year IS NULL, b.year", "title": "b.title", "author": "b.authors IS NULL, b.authors",
             "size": "b.size DESC", "pages": "b.pages DESC"}[a.sort]
    sql = f"SELECT b.* FROM books b {'WHERE ' + ' AND '.join(where) if where else ''} ORDER BY {order} LIMIT ?"
    return db.execute(sql, params + [a.limit]).fetchall()


def cmd_search(db, lib, a):
    rows = search(db, a)
    if not rows:
        print("No matches."); return
    w = shutil_width()
    for i, r in enumerate(rows, 1):
        yr = r["year_text"] or (str(r["year"]) if r["year"] is not None else "----")
        mark = "" if r["downloaded"] else "  [not downloaded yet]"
        print(f"{i:>3}. [{r['id']}] {r['title']}")
        info = " | ".join(filter(None, [r["authors"], yr, r["publisher"], r["doc_type"], r["language"],
                                        f"{r['pages']} p" if r["pages"] else None, r["tags"] or None]))
        print(f"      {info}{mark}")
        print(f"      {r['topic']}{' › ' + r['subtopic'] if r['subtopic'] else ''}"[:w])
        if a.paths:
            print(f"      {os.path.join(lib, r['local_path'])}")
    print(f"\n{len(rows)} result(s){' (limit reached, use --limit)' if len(rows) == a.limit else ''}")
    if a.open:
        r = rows[a.open - 1]
        open_file(lib, r)


def shutil_width():
    try:
        return os.get_terminal_size().columns - 2
    except OSError:
        return 160


def open_file(lib, r):
    p = os.path.join(lib, r["local_path"])
    if not os.path.exists(long_path(p)):
        print(f"File not downloaded yet: {p}"); return
    print(f"Opening {p}")
    if os.name == "nt":
        os.startfile(p if len(p) < 250 else long_path(p))
    else:
        import subprocess
        subprocess.run(["xdg-open" if sys.platform.startswith("linux") else "open", p])


def cmd_open(db, lib, a):
    r = db.execute("SELECT * FROM books WHERE id=?", (a.id,)).fetchone()
    if not r:
        print("No such id."); return
    open_file(lib, r)


def cmd_list(db, lib, a, kind):
    col = {"authors": "authors", "topics": "topic", "publishers": "publisher", "years": "year",
           "languages": "language", "types": "doc_type", "traditions": "tradition"}[kind]
    if kind == "authors":
        c = Counter()
        for (v,) in db.execute("SELECT authors FROM books WHERE authors IS NOT NULL"):
            for x in v.split("; "):
                c[x] += 1
        items = c.most_common()
    else:
        items = db.execute(f"SELECT {col}, COUNT(*) FROM books WHERE {col} IS NOT NULL GROUP BY {col} "
                           f"ORDER BY {'1' if kind == 'years' else '2 DESC'}").fetchall()
    if a.like:
        items = [(k, n) for k, n in items if a.like.lower() in str(k).lower()]
    for k, n in items[: a.limit]:
        print(f"{n:>6}  {k}")
    print(f"\n{len(items)} distinct value(s)")


def cmd_stats(db, lib, a):
    total, size, dl = db.execute("SELECT COUNT(*), COALESCE(SUM(size),0), SUM(downloaded) FROM books").fetchone()
    print(f"\n  Entries      : {total}   (downloaded {dl or 0}),  {size / 1024 ** 3:.1f} GB")
    for label, col in (("with author", "authors"), ("with year", "year"), ("with publisher", "publisher"),
                       ("with pages", "pages")):
        n = db.execute(f"SELECT COUNT(*) FROM books WHERE {col} IS NOT NULL").fetchone()[0]
        print(f"  {label:<13}: {n}  ({100 * n / max(total, 1):.0f}%)")
    for title, col in (("Types", "doc_type"), ("Traditions", "tradition"), ("Languages", "language")):
        top = db.execute(f"SELECT {col}, COUNT(*) FROM books GROUP BY {col} ORDER BY 2 DESC LIMIT 8").fetchall()
        print(f"  {title:<13}: " + ", ".join(f"{k} {n}" for k, n in top))


def cmd_export(db, lib, a):
    rows = db.execute("SELECT id,title,authors,year_text,year,publisher,doc_type,tradition,language,tags,topic,"
                      "subtopic,pages,size,ext,downloaded,local_path FROM books ORDER BY topic, title").fetchall()
    headers = rows[0].keys() if rows else []
    if a.csv:
        with open(a.csv, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh); w.writerow(headers); w.writerows([tuple(r) for r in rows])
        print(f"Wrote {len(rows)} rows to {a.csv}")
    if a.xlsx:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook(); ws = wb.active; ws.title = "Catalogue"
        ws.append(list(headers))
        for c in ws[1]:
            c.font = Font(bold=True)
        illegal = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
        for i, r in enumerate(rows, start=2):
            ws.append([illegal.sub("", v) if isinstance(v, str) else v for v in r])
            ws.cell(row=i, column=2).hyperlink = os.path.join(lib, r["local_path"])
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "C2"
        for col, width in zip("ABCDEFGHIJKLMNOPQ", (7, 60, 28, 8, 7, 24, 14, 18, 12, 18, 30, 40, 7, 12, 6, 6, 60)):
            ws.column_dimensions[col].width = width
        wb.save(a.xlsx)
        print(f"Wrote {len(rows)} rows to {a.xlsx}")


# ----------------------------------------------------------------------------- browse folders (hard links)
def safe(s, n=120):
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(s)).strip(" .")
    return (s[:n].strip(" .") or "_")


def cmd_views(db, lib, a):
    kinds = [k.strip() for k in a.by.split(",")]
    base = os.path.join(lib, "_Browse")
    made = skipped = 0
    rows = db.execute("SELECT * FROM books WHERE downloaded=1").fetchall()
    for kind in kinds:
        for r in rows:
            if kind == "author":
                keys = (r["authors"] or "_Unknown author").split("; ")
                dirs = [os.path.join("By Author", safe(k[0].upper() if k[0].isalpha() else "#"), safe(k)) for k in keys]
            elif kind == "year":
                if r["year"] is None:
                    dirs = [os.path.join("By Year", "_Unknown year")]
                else:
                    dec = f"{(r['year'] // 10) * 10}s"
                    dirs = [os.path.join("By Year", safe(dec), safe(r["year"]))]
            elif kind == "publisher":
                dirs = [os.path.join("By Publisher", safe(r["publisher"] or "_Unknown publisher"))]
            elif kind == "language":
                dirs = [os.path.join("By Language", safe(r["language"] or "_Unknown"))]
            elif kind == "type":
                dirs = [os.path.join("By Type", safe(r["doc_type"] or "_Unknown"))]
            elif kind == "topic":
                dirs = [os.path.join("By Topic", safe(r["topic"] or "_Unknown"))]
            elif kind == "tradition":
                dirs = [os.path.join("By Tradition", safe(r["tradition"] or "_Unknown"))]
            else:
                sys.exit(f"unknown view: {kind}")
            src = long_path(os.path.join(lib, r["local_path"]))
            if not os.path.exists(src):
                continue
            fname = os.path.basename(r["local_path"])
            for d in dirs:
                dst_dir = long_path(os.path.join(base, d))
                os.makedirs(dst_dir, exist_ok=True)
                dst = os.path.join(dst_dir, fname)
                if os.path.exists(dst):
                    skipped += 1; continue
                try:
                    os.link(src, dst); made += 1
                except OSError as e:
                    print(f"  could not link {fname}: {e}")
                    if a.stop_on_error:
                        return
        print(f"  view '{kind}' done")
    print(f"Created {made} hard links ({skipped} already existed) under {base}")
    print("Hard links share the same data as the original file: no extra space is used, and deleting a link "
          "does not delete the book.")


# ----------------------------------------------------------------------------- offline HTML catalogue
HTML_TEMPLATE = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Astro Library Catalogue</title>
<style>
:root{--bg:#fbfaf7;--fg:#1f1d1a;--mute:#6b665e;--line:#e4e0d8;--acc:#7a3e9d;--chip:#f1ecf5;--row:#fff}
@media (prefers-color-scheme:dark){:root{--bg:#16151a;--fg:#ebe8f0;--mute:#9a96a3;--line:#2c2a33;--acc:#c79bea;--chip:#2a2433;--row:#1c1b21}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:14px/1.45 system-ui,Segoe UI,sans-serif}
header{padding:18px 20px 8px}h1{margin:0;font-size:22px}h1 small{font-weight:400;color:var(--mute);font-size:13px;margin-left:8px}
.f{display:grid;grid-template-columns:repeat(auto-fill,minmax(170px,1fr));gap:8px;padding:8px 20px;position:sticky;top:0;background:var(--bg);border-bottom:1px solid var(--line);z-index:2}
.f input,.f select{width:100%;padding:7px 8px;border:1px solid var(--line);border-radius:6px;background:var(--row);color:var(--fg);font:inherit}
.f .q{grid-column:span 2}.bar{display:flex;gap:12px;align-items:center;padding:8px 20px;color:var(--mute);flex-wrap:wrap}
.bar button{border:1px solid var(--line);background:var(--row);color:var(--fg);padding:5px 10px;border-radius:6px;cursor:pointer}
table{width:100%;border-collapse:collapse}th,td{padding:7px 10px;border-bottom:1px solid var(--line);text-align:left;vertical-align:top}
th{position:sticky;top:0;cursor:pointer;color:var(--mute);font-weight:600;font-size:12px;user-select:none;background:var(--bg)}
td.t a{color:var(--fg);text-decoration:none;font-weight:600}td.t a:hover{color:var(--acc);text-decoration:underline}
td .s{color:var(--mute);font-size:12px}.chip{display:inline-block;background:var(--chip);color:var(--acc);border-radius:10px;padding:0 7px;font-size:11px;margin:2px 3px 0 0}
.nd{opacity:.55}.wrap{padding:0 10px 40px;overflow-x:auto}.y{white-space:nowrap}
@media (max-width:700px){.hide-m{display:none}.f .q{grid-column:span 1}}
</style></head><body>
<header><h1>Astro Library Catalogue<small id="sum"></small></h1></header>
<div class="f">
<input class="q" id="q" placeholder="Search title, author, topic, publisher, file name…" autofocus>
<input id="au" placeholder="Author"><input id="pu" placeholder="Publisher">
<input id="y1" type="number" placeholder="Year from"><input id="y2" type="number" placeholder="Year to">
<select id="to"><option value="">All topics</option></select>
<select id="ty"><option value="">All types</option></select>
<select id="tr"><option value="">All traditions</option></select>
<select id="la"><option value="">All languages</option></select>
<select id="tg"><option value="">Any tag</option></select>
</div>
<div class="bar"><span id="cnt"></span><button id="prev">‹ Prev</button><span id="pg"></span><button id="next">Next ›</button><button id="reset">Reset filters</button><label><input type="checkbox" id="dl"> downloaded only</label></div>
<div class="wrap"><table><thead><tr><th data-k="1">Title</th><th data-k="2">Author</th><th data-k="3" class="y">Year</th><th data-k="4" class="hide-m">Publisher</th><th data-k="7" class="hide-m">Topic</th><th data-k="5" class="hide-m">Type</th><th data-k="13" class="hide-m">Pages</th></tr></thead><tbody id="tb"></tbody></table></div>
<script>
const D=__DATA__;const GEN="__GEN__";
// row: 0 id,1 title,2 authors,3 year,4 publisher,5 type,6 tradition,7 topic,8 subtopic,9 language,10 tags,11 path,12 downloaded,13 pages,14 sizeMB,15 yeartext
const $=id=>document.getElementById(id);const PS=100;let page=0,sortK=null,sortD=1,res=D;
D.forEach(r=>r.push(((r[1]||'')+' '+(r[2]||'')+' '+(r[4]||'')+' '+(r[7]||'')+' '+(r[8]||'')+' '+(r[10]||'')+' '+r[11]).toLowerCase()));
function fill(id,idx,split){const c={};D.forEach(r=>{const v=r[idx];if(!v)return;(split?v.split(', '):[v]).forEach(x=>c[x]=(c[x]||0)+1)});
Object.keys(c).sort((a,b)=>c[b]-c[a]).forEach(k=>{const o=document.createElement('option');o.value=k;o.textContent=k+' ('+c[k]+')';$(id).appendChild(o)})}
fill('to',7);fill('ty',5);fill('tr',6);fill('la',9);fill('tg',10,true);
const esc=s=>(s==null?'':String(s)).replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
const href=p=>p.split('/').map(encodeURIComponent).join('/');
function run(){const q=$('q').value.toLowerCase().split(/\s+/).filter(Boolean),au=$('au').value.toLowerCase(),pu=$('pu').value.toLowerCase(),
y1=parseInt($('y1').value),y2=parseInt($('y2').value),to=$('to').value,ty=$('ty').value,tr=$('tr').value,la=$('la').value,tg=$('tg').value,dl=$('dl').checked;
res=D.filter(r=>(!q.length||q.every(w=>r[16].includes(w)))&&(!au||(r[2]||'').toLowerCase().includes(au))&&(!pu||(r[4]||'').toLowerCase().includes(pu))
&&(isNaN(y1)||(r[3]!=null&&r[3]>=y1))&&(isNaN(y2)||(r[3]!=null&&r[3]<=y2))&&(!to||r[7]===to)&&(!ty||r[5]===ty)&&(!tr||r[6]===tr)&&(!la||r[9]===la)
&&(!tg||(r[10]||'').split(', ').includes(tg))&&(!dl||r[12]));
if(sortK!=null)res.sort((a,b)=>{const x=a[sortK],y=b[sortK];if(x==null)return 1;if(y==null)return -1;return (x>y?1:x<y?-1:0)*sortD});
page=0;draw()}
function draw(){const n=res.length,pages=Math.max(1,Math.ceil(n/PS));$('cnt').textContent=n.toLocaleString()+' match'+(n==1?'':'es');$('pg').textContent='page '+(page+1)+' / '+pages;
$('tb').innerHTML=res.slice(page*PS,page*PS+PS).map(r=>'<tr class="'+(r[12]?'':'nd')+'"><td class="t"><a href="'+href(r[11])+'" target="_blank" title="'+esc(r[11])+'">'+esc(r[1])+'</a>'+
(r[10]?'<div>'+r[10].split(', ').map(t=>'<span class="chip">'+esc(t)+'</span>').join('')+'</div>':'')+'<div class="s">'+esc(r[9])+' · '+esc(r[6])+(r[12]?'':' · not downloaded yet')+'</div></td><td>'+esc(r[2])+'</td><td class="y">'+esc(r[15]||'')+'</td><td class="hide-m">'+esc(r[4])+'</td><td class="hide-m">'+esc(r[7])+(r[8]?'<div class="s">'+esc(r[8])+'</div>':'')+'</td><td class="hide-m">'+esc(r[5])+'</td><td class="hide-m">'+(r[13]||'')+'</td></tr>').join('');window.scrollTo(0,0)}
let tm;document.querySelectorAll('.f input,.f select,#dl').forEach(e=>e.addEventListener('input',()=>{clearTimeout(tm);tm=setTimeout(run,180)}));
$('prev').onclick=()=>{if(page>0){page--;draw()}};$('next').onclick=()=>{if((page+1)*PS<res.length){page++;draw()}};
$('reset').onclick=()=>{document.querySelectorAll('.f input').forEach(e=>e.value='');document.querySelectorAll('.f select').forEach(e=>e.value='');$('dl').checked=false;run()};
document.querySelectorAll('th').forEach(th=>th.onclick=()=>{const k=+th.dataset.k;sortD=sortK===k?-sortD:1;sortK=k;run()});
$('sum').textContent=D.length.toLocaleString()+' items · generated '+GEN;run();
</script></body></html>"""


def cmd_html(db, lib, a):
    rows = db.execute("SELECT * FROM books ORDER BY title COLLATE NOCASE").fetchall()
    data = [[r["id"], r["title"], r["authors"], r["year"], r["publisher"], r["doc_type"], r["tradition"], r["topic"],
             r["subtopic"], r["language"], r["tags"] or None, r["local_path"], r["downloaded"], r["pages"],
             round((r["size"] or 0) / 1048576, 1), r["year_text"]] for r in rows]
    js = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    out = os.path.join(lib, "Astro Library Catalogue.html")
    with open(out, "w", encoding="utf-8") as fh:
        fh.write(HTML_TEMPLATE.replace("__DATA__", js).replace("__GEN__", datetime.now().strftime("%Y-%m-%d %H:%M")))
    print(f"Offline catalogue written: {out}")


# ----------------------------------------------------------------------------- CLI
def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Catalogue and search the astrology library.",
                                 formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    ap.add_argument("--lib", default=DEFAULT_LIB, help=f"library folder (default {DEFAULT_LIB})")
    sub = ap.add_subparsers(dest="cmd", required=True)

    b = sub.add_parser("build"); b.add_argument("--no-pdf-meta", action="store_true")

    s = sub.add_parser("search")
    s.add_argument("text", nargs="*", help="free text (all words must match; prefix match; \"exact phrase\")")
    s.add_argument("--author"); s.add_argument("--title"); s.add_argument("--topic"); s.add_argument("--publisher")
    s.add_argument("--year", type=int); s.add_argument("--from", dest="year_from", type=int)
    s.add_argument("--to", dest="year_to", type=int); s.add_argument("--lang"); s.add_argument("--type")
    s.add_argument("--tradition"); s.add_argument("--tag"); s.add_argument("--ext")
    s.add_argument("--downloaded", action="store_true", help="only files already on disk")
    s.add_argument("--sort", choices=["year", "title", "author", "size", "pages"], default="title")
    s.add_argument("--limit", type=int, default=50); s.add_argument("--paths", action="store_true")
    s.add_argument("--open", type=int, metavar="N", help="open result number N")

    for k in ("authors", "topics", "publishers", "years", "languages", "types", "traditions"):
        p = sub.add_parser(k); p.add_argument("--like"); p.add_argument("--limit", type=int, default=200)
    sub.add_parser("stats")
    o = sub.add_parser("open"); o.add_argument("id", type=int)
    e = sub.add_parser("export"); e.add_argument("--csv"); e.add_argument("--xlsx")
    sub.add_parser("html")
    v = sub.add_parser("views"); v.add_argument("--by", default="author,year,publisher")
    v.add_argument("--stop-on-error", action="store_true")
    a = ap.parse_args()

    lib = os.path.abspath(a.lib)
    db = open_db(lib)
    if a.cmd != "build" and not db.execute("SELECT name FROM sqlite_master WHERE name='books'").fetchone():
        sys.exit("Catalogue not built yet. Run:  python astro_library.py build")
    if a.cmd == "build":
        cmd_build(db, lib, a)
    elif a.cmd == "search":
        cmd_search(db, lib, a)
    elif a.cmd in ("authors", "topics", "publishers", "years", "languages", "types", "traditions"):
        cmd_list(db, lib, a, a.cmd)
    elif a.cmd == "stats":
        cmd_stats(db, lib, a)
    elif a.cmd == "open":
        cmd_open(db, lib, a)
    elif a.cmd == "export":
        if not (a.csv or a.xlsx):
            sys.exit("give --csv FILE and/or --xlsx FILE")
        cmd_export(db, lib, a)
    elif a.cmd == "html":
        cmd_html(db, lib, a)
    elif a.cmd == "views":
        cmd_views(db, lib, a)


if __name__ == "__main__":
    main()
